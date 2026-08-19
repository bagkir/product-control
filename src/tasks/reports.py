import asyncio
import os
import tempfile
from datetime import UTC, datetime, timedelta

from openpyxl import Workbook

from src.celery_app import celery_app
from src.core.celery_database import celery_db_session
from src.data.repositories import (
    WebhookDeliveryRepository,
    WebhookSubscriptionRepository,
)
from src.data.repositories.batch_repository import BatchRepository
from src.domain.services.webhook_service import WebhookService
from src.storage.minio_service import MinIOService
from src.utils.pdf_generator import build_batch_pdf_report

REPORTS_BUCKET = "reports"


@celery_app.task(bind=True, max_retries=3, name="tasks.generate_batch_report")
def generate_batch_report(
    self,
    batch_id: int,
    format: str = "excel",
    user_email: str | None = None,
) -> dict:
    return asyncio.run(_generate_batch_report_async(self, batch_id, format, user_email))


async def _generate_batch_report_async(task, batch_id: int, format: str) -> dict:
    async with celery_db_session() as session:
        webhook_service = WebhookService(
            subscription_repository=WebhookSubscriptionRepository(session),
            delivery_repository=WebhookDeliveryRepository(session),
        )
        batch_repository = BatchRepository(session)
        batch = await batch_repository.get_full_for_report(batch_id)
        if batch is None:
            raise ValueError(f"Batch {batch_id} not found")

        task.update_state(
            state="PROGRESS", meta={"current": 1, "total": 3, "progress": 33.0}
        )

        if format == "excel":
            file_path, file_name = _build_excel_report(batch)
        elif format == "pdf":
            file_path, file_name = build_batch_pdf_report(batch)
        else:
            raise ValueError(f"Unsupported report format: {format}")

        task.update_state(
            state="PROGRESS", meta={"current": 2, "total": 3, "progress": 66.0}
        )

        try:
            file_size = os.path.getsize(file_path)

            storage = MinIOService()
            file_url = await asyncio.to_thread(
                storage.upload_file,
                bucket=REPORTS_BUCKET,
                file_path=file_path,
                object_name=file_name,
                expires_days=7,
            )
        finally:
            os.remove(file_path)

        expires_at = datetime.now(UTC) + timedelta(days=7)

        task.update_state(
            state="PROGRESS", meta={"current": 3, "total": 3, "progress": 100.0}
        )

        await webhook_service.publish_event(
            event_type="report_generated",
            data={
                "batch_id": batch_id,
                "report_type": format,
                "file_url": file_url,
                "expires_at": expires_at.isoformat(),
            },
        )

        return {
            "success": True,
            "file_url": file_url,
            "file_name": file_name,
            "file_size": file_size,
            "expires_at": expires_at.isoformat(),
        }


def _build_excel_report(batch) -> tuple[str, str]:
    """
    Строит xlsx с тремя листами (см. ТЗ) и сохраняет во временный файл.
    Возвращает (путь_к_файлу, имя_файла).
    """
    wb = Workbook()

    # --- Лист 1: Информация о партии ---
    ws_info = wb.active
    ws_info.title = "Информация о партии"
    info_rows = [
        ("Номер партии", batch.batch_number),
        ("Дата партии", batch.batch_date.isoformat()),
        ("Статус", "Закрыта" if batch.is_closed else "Открыта"),
        ("Рабочий центр", batch.work_center.name if batch.work_center else ""),
        ("Смена", batch.shift),
        ("Бригада", batch.team),
        ("Номенклатура", batch.nomenclature),
        ("Начало смены", batch.shift_start.strftime("%Y-%m-%d %H:%M:%S")),
        ("Окончание смены", batch.shift_end.strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row in info_rows:
        ws_info.append(row)

    # --- Лист 2: Продукция ---
    ws_products = wb.create_sheet("Продукция")
    ws_products.append(["ID", "Уникальный код", "Аггрегирована", "Дата аггрегации"])
    for product in batch.products:
        ws_products.append(
            [
                product.id,
                product.unique_code,
                "Да" if product.is_aggregated else "Нет",
                (
                    product.aggregated_at.strftime("%Y-%m-%d %H:%M:%S")
                    if product.aggregated_at
                    else "-"
                ),
            ]
        )

    # --- Лист 3: Статистика ---
    total = len(batch.products)
    aggregated = sum(1 for p in batch.products if p.is_aggregated)
    remaining = total - aggregated
    rate = (aggregated / total * 100) if total else 0.0

    elapsed_hours = _elapsed_hours(batch)
    speed = (aggregated / elapsed_hours) if elapsed_hours > 0 else 0.0

    ws_stats = wb.create_sheet("Статистика")
    ws_stats.append(["Всего продукции", total])
    ws_stats.append(["Аггрегировано", aggregated])
    ws_stats.append(["Осталось", remaining])
    ws_stats.append(["Процент выполнения", f"{rate:.1f}%"])
    ws_stats.append(["Средняя скорость", f"{speed:.1f} ед/час"])

    file_name = f"batch_{batch.id}_report.xlsx"
    file_path = os.path.join(tempfile.gettempdir(), file_name)
    wb.save(file_path)

    return file_path, file_name


def _elapsed_hours(batch) -> float:
    """Часы с начала смены до текущего момента (или до конца смены, если она уже прошла)."""
    now = datetime.now(UTC)
    shift_start = (
        batch.shift_start.replace(tzinfo=UTC)
        if batch.shift_start.tzinfo is None
        else batch.shift_start
    )
    shift_end = (
        batch.shift_end.replace(tzinfo=UTC)
        if batch.shift_end.tzinfo is None
        else batch.shift_end
    )

    reference = min(now, shift_end)
    delta_hours = (reference - shift_start).total_seconds() / 3600
    return max(delta_hours, 0.0)
