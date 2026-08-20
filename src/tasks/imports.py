import asyncio
import csv
import os
import tempfile

import aiofiles
import httpx
from celery import Task
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError

from src.api.v1.schemas.batch import BatchCreate
from src.celery_app import celery_app
from src.core.celery_database import celery_db_session
from src.data.repositories import (
    WebhookDeliveryRepository,
    WebhookSubscriptionRepository,
)
from src.data.repositories.batch_repository import BatchRepository
from src.data.repositories.work_center_repository import WorkCenterRepository
from src.domain.services.webhook_service import WebhookService

IMPORTS_BUCKET = "imports"


@celery_app.task(
    bind=True,
    max_retries=1,
    name="tasks.import_batches_from_file",
)
def import_batches_from_file(
    self: Task,
    file_url: str,
) -> dict:
    return asyncio.run(
        _import_batches_from_file_async(
            self,
            file_url,
        )
    )


async def _import_batches_from_file_async(task: Task, file_url: str) -> dict:
    file_path = None

    try:
        file_path = await _download_file(file_url)

        rows = _read_rows(file_path)

        total = len(rows)
        created = 0
        skipped = 0
        errors: list[dict] = []

        async with celery_db_session() as session:
            webhook_service = WebhookService(
                subscription_repository=WebhookSubscriptionRepository(session),
                delivery_repository=WebhookDeliveryRepository(session),
            )
            batch_repository = BatchRepository(session)
            work_center_repository = WorkCenterRepository(session)

            task.update_state(
                state="PROGRESS",
                meta={
                    "current": 0,
                    "total": total,
                    "created": 0,
                    "skipped": 0,
                },
            )

            for index, row in enumerate(rows, start=2):
                try:
                    data = BatchCreate.model_validate(row)

                    if await batch_repository.exists_by_number_and_date(
                        data.batch_number, data.batch_date
                    ):
                        skipped += 1
                        errors.append(
                            {"row": index, "error": "Duplicate batch number and date"}
                        )
                        continue
                    work_center = await work_center_repository.get_by_identifier(
                        data.work_center_identifier
                    )
                    if work_center is None:
                        work_center = await work_center_repository.create(
                            identifier=data.work_center_identifier,
                            name=data.work_center,
                        )
                    await batch_repository.create(
                        is_closed=data.is_closed,
                        task_description=data.task_description,
                        work_center_id=work_center.id,
                        shift=data.shift,
                        team=data.team,
                        batch_number=data.batch_number,
                        batch_date=data.batch_date,
                        nomenclature=data.nomenclature,
                        ekn_code=data.ekn_code,
                        shift_start=data.shift_start,
                        shift_end=data.shift_end,
                    )

                    await session.commit()
                    created += 1

                except ValidationError as exc:
                    await session.rollback()

                    errors.append(
                        {
                            "row": index,
                            "error": _format_validation_error(exc),
                        }
                    )

                except IntegrityError:
                    await session.rollback()
                    skipped += 1
                    errors.append(
                        {
                            "row": index,
                            "error": "Integrity constraint violation (likely duplicate)",
                        }
                    )

                except Exception as exc:  # noqa: BLE001
                    await session.rollback()

                    errors.append(
                        {
                            "row": index,
                            "error": str(exc),
                        }
                    )

                task.update_state(
                    state="PROGRESS",
                    meta={
                        "current": index - 1,
                        "total": total,
                        "created": created,
                        "skipped": skipped,
                    },
                )
            await webhook_service.publish_event(
                event_type="import_completed",
                data={
                    "total_rows": total,
                    "created": created,
                    "skipped": skipped,
                    "errors": errors,
                },
            )
        return {
            "success": len(errors) == 0,
            "total_rows": total,
            "created": created,
            "skipped": skipped,
            "errors": errors,
        }

    except Exception as exc:  # noqa: BLE001
        raise task.retry(exc=exc)

    finally:
        if file_path and os.path.exists(file_path):
            os.remove(file_path)


async def _download_file(file_url: str) -> str:
    extension = _get_extension(file_url)

    with tempfile.NamedTemporaryFile(suffix=extension, delete=False) as tmp:
        temp_path = tmp.name

    async with httpx.AsyncClient() as client:
        response = await client.get(file_url)
        response.raise_for_status()
        async with aiofiles.open(temp_path, "wb") as f:
            await f.write(response.content)

    return temp_path


def _read_rows(file_path: str) -> list[dict]:
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".csv":
        return _read_csv(file_path)

    if extension == ".xlsx":
        return _read_xlsx(file_path)

    raise ValueError(f"Unsupported file format: {extension}")


def _read_xlsx(file_path: str) -> list[dict]:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    rows = worksheet.iter_rows(values_only=True)

    headers = next(rows)

    result = []

    for values in rows:
        if not any(value is not None for value in values):
            continue

        row = dict(zip(headers, values))

        result.append(row)

    workbook.close()

    return result


def _read_csv(file_path: str) -> list[dict]:
    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        return list(reader)


def _get_extension(file_url: str) -> str:
    path = file_url.split("?", maxsplit=1)[0]

    extension = os.path.splitext(path)[1].lower()

    return extension or ".xlsx"


def _format_validation_error(exc: ValidationError) -> str:
    return "; ".join(f"{error['loc'][0]}: {error['msg']}" for error in exc.errors())
